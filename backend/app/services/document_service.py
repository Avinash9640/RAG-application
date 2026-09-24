from pathlib import Path
import re
from uuid import uuid4

from fastapi import HTTPException, UploadFile, status
from docx import Document as DocxDocument
from pypdf import PdfReader

from app.models.document import Document
from app.models.document_chunk import DocumentChunk
from app.services.embedding_service import embed_texts


UPLOAD_DIRECTORY = Path(__file__).resolve().parents[2] / "storage" / "uploads"
ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt"}
MAX_FILE_SIZE_BYTES = 25 * 1024 * 1024
DEFAULT_CHUNK_SIZE = 1000
DEFAULT_CHUNK_OVERLAP = 200


def save_document(
    upload: UploadFile,
    uploaded_by: str,
    chunk_strategy: str,
    chunk_size: int,
    chunk_overlap: int,
) -> Document:
    filename = Path(upload.filename or "").name
    extension = Path(filename).suffix.lower()
    if not filename or extension not in ALLOWED_EXTENSIONS:
        raise HTTPException(status_code=400, detail="Only PDF, DOCX, and TXT files are supported.")

    UPLOAD_DIRECTORY.mkdir(parents=True, exist_ok=True)
    stored_path = UPLOAD_DIRECTORY / f"{uuid4()}{extension}"
    bytes_written = 0
    try:
        with stored_path.open("wb") as destination:
            while chunk := upload.file.read(1024 * 1024):
                bytes_written += len(chunk)
                if bytes_written > MAX_FILE_SIZE_BYTES:
                    raise HTTPException(status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE, detail="Document must be 25 MB or smaller.")
                destination.write(chunk)
    except Exception:
        stored_path.unlink(missing_ok=True)
        raise
    finally:
        upload.file.close()

    return Document(id=str(uuid4()), filename=filename,
                    file_type=extension.removeprefix(".").upper(),
                    storage_path=str(stored_path), file_size=bytes_written,
                    status="Uploaded", uploaded_by=uploaded_by,
                    allowed_roles=["User", "Admin"],
                    chunk_strategy=chunk_strategy, chunk_size=chunk_size,
                    chunk_overlap=chunk_overlap)


def remove_stored_document(document: Document) -> None:
    Path(document.storage_path).unlink(missing_ok=True)


def remove_stored_file(storage_path: str) -> None:
    Path(storage_path).unlink(missing_ok=True)


def extract_pages(document: Document) -> list[tuple[int, str]]:
    path = Path(document.storage_path)
    if document.file_type == "PDF":
        reader = PdfReader(path)
        return [(number, page.extract_text() or "") for number, page in enumerate(reader.pages, start=1)]
    if document.file_type == "DOCX":
        paragraphs = DocxDocument(path).paragraphs
        return [(1, "\n".join(paragraph.text for paragraph in paragraphs))]
    if document.file_type == "TXT":
        return [(1, path.read_text(encoding="utf-8", errors="replace"))]
    raise HTTPException(status_code=400, detail="Unsupported document type.")


def split_text(text: str, chunk_size: int = DEFAULT_CHUNK_SIZE,
               overlap: int = DEFAULT_CHUNK_OVERLAP,
               strategy: str = "recursive") -> list[str]:
    cleaned = re.sub(r"[ \t]+", " ", text.replace("\r\n", "\n"))
    cleaned = re.sub(r"\n{3,}", "\n\n", cleaned).strip()
    if not cleaned:
        return []

    normalized = " ".join(cleaned.split())
    if strategy == "fixed":
        return _window_chunks(normalized, chunk_size, overlap, [])
    if strategy == "sentence":
        return _unit_chunks(
            re.split(r"(?<=[.!?])\s+", normalized), chunk_size, overlap
        )
    if strategy == "paragraph":
        paragraphs = [part.strip() for part in re.split(r"\n\s*\n", cleaned) if part.strip()]
        return _unit_chunks(paragraphs, chunk_size, overlap)
    return _window_chunks(
        cleaned, chunk_size, overlap, ["\n\n", "\n", ". ", "; ", ", ", " "]
    )


def _window_chunks(
    text: str, chunk_size: int, overlap: int, separators: list[str]
) -> list[str]:
    chunks: list[str] = []
    start = 0
    while start < len(text):
        end = min(start + chunk_size, len(text))
        if end < len(text):
            for separator in separators:
                boundary = text.rfind(separator, start + 1, end + 1)
                if boundary > start:
                    end = boundary + (len(separator) if separator.strip() else 0)
                    break
        content = text[start:end].strip()
        if content:
            chunks.append(content)
        if end >= len(text):
            break
        start = max(end - overlap, start + 1)
    return chunks


def _unit_chunks(units: list[str], chunk_size: int, overlap: int) -> list[str]:
    chunks: list[str] = []
    current = ""
    for unit in (item.strip() for item in units if item.strip()):
        if len(unit) > chunk_size:
            if current:
                chunks.append(current)
                current = ""
            chunks.extend(_window_chunks(unit, chunk_size, overlap, [" "]))
            continue
        candidate = f"{current} {unit}".strip()
        if current and len(candidate) > chunk_size:
            chunks.append(current)
            prefix = current[-overlap:].lstrip() if overlap else ""
            current = f"{prefix} {unit}".strip()
            if len(current) > chunk_size:
                current = unit
        else:
            current = candidate
    if current:
        chunks.append(current)
    return chunks


def create_document_chunks(document: Document) -> list[DocumentChunk]:
    chunk_values = []
    for page_number, text in extract_pages(document):
        for chunk_index, content in enumerate(split_text(
            text, document.chunk_size, document.chunk_overlap,
            document.chunk_strategy,
        )):
            chunk_values.append((page_number, chunk_index, content))
    if not chunk_values:
        raise HTTPException(status_code=422, detail="No extractable text was found in this document.")

    embeddings = embed_texts([content for _, _, content in chunk_values])
    chunks = []
    for (page_number, chunk_index, content), embedding in zip(chunk_values, embeddings):
        chunks.append(DocumentChunk(
                id=str(uuid4()),
                document_id=document.id,
                page_number=page_number,
                chunk_index=chunk_index,
                content=content,
                chunk_size=document.chunk_size,
                overlap=document.chunk_overlap,
                embedding=embedding,
                chunk_metadata={"filename": document.filename},
            ))
    return chunks
